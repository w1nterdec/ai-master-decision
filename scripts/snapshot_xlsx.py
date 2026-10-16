#!/usr/bin/env python3
"""对 Phase 3.7C 研究 Excel 做全量快照，用于修改前后逐格差异比对。

同时抓取两种视图：
    * formula_view（data_only=False）：保留公式原文，用于确认公式是否被改动；
    * value_view  （data_only=True） ：公式的计算结果，用于确认数值是否被改动。

用法：
    python3 scripts/snapshot_xlsx.py <xlsx路径> <输出json路径>
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl


def snapshot(path: str) -> dict:
    out: dict = {"file": path, "sheets": {}}

    wb_f = openpyxl.load_workbook(path, data_only=False)
    out["sheetnames"] = list(wb_f.sheetnames)
    out["sheet_order"] = {n: i for i, n in enumerate(wb_f.sheetnames, 1)}
    out["sheet_props"] = {
        n: {
            "max_row": wb_f[n].max_row,
            "max_column": wb_f[n].max_column,
            "dimensions": wb_f[n].dimensions,
            "sheet_state": wb_f[n].sheet_state,
            "freeze": str(wb_f[n].freeze_panes),
        }
        for n in wb_f.sheetnames
    }

    for name in wb_f.sheetnames:
        ws = wb_f[name]
        cells = {}
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None and c.data_type == "n":
                    continue
                if v is None:
                    continue
                cells[c.coordinate] = {
                    "v": str(v),
                    "t": c.data_type,
                    "f": "F" if isinstance(v, str) and v.startswith("=") else "",
                }
        out["sheets"][name] = {"formula_view": cells}
    wb_f.close()

    try:
        wb_v = openpyxl.load_workbook(path, data_only=True)
        for name in wb_v.sheetnames:
            ws = wb_v[name]
            cells = {}
            for row in ws.iter_rows():
                for c in row:
                    if c.value is None:
                        continue
                    cells[c.coordinate] = str(c.value)
            out["sheets"][name]["value_view"] = cells
        wb_v.close()
    except Exception as exc:  # noqa: BLE001
        out["value_view_error"] = str(exc)

    return out


def main() -> int:
    if len(sys.argv) != 3:
        print("用法: python3 scripts/snapshot_xlsx.py <xlsx> <out.json>")
        return 1
    data = snapshot(sys.argv[1])
    Path(sys.argv[2]).write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    n = sum(len(s.get("formula_view", {})) for s in data["sheets"].values())
    print(f"[OK] 快照完成：{len(data['sheetnames'])} 张表 / {n} 个非空单元格 -> {sys.argv[2]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

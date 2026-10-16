"""测试公共夹具与常量。

PHASE37C_ROWS 为 Phase 3.7C Excel「三级指标树」的 18 条三级指标原始权重
（一级权重 / 二级组内权重 / 三级组内权重，单位：%）。
golden 测试会现场读取 Excel 校验该常量未被篡改。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from decision_system.engine.weights import WeightTree, build_tree_from_rows

ROOT = Path(__file__).resolve().parents[1]
RAW_XLSX = ROOT / "data" / "raw" / "Phase3_7C_情景覆盖率门槛修复版.xlsx"
RAW_JSON = ROOT / "data" / "raw" / "Phase3_7C_实际回归测试结果.json"

PHASE37C_ROWS: list[dict[str, Any]] = [
    {"level1": "职业发展", "level2": "AI工程就业", "level3": "岗位与实习机会", "w1": 35, "w2": 50, "w3": 60},
    {"level1": "职业发展", "level2": "AI工程就业", "level3": "产业合作可参与性", "w1": 35, "w2": 50, "w3": 40},
    {"level1": "职业发展", "level2": "长期成长", "level3": "Agent/LLM能力积累", "w1": 35, "w2": 50, "w3": 60},
    {"level1": "职业发展", "level2": "长期成长", "level3": "具身/机器人发展空间", "w1": 35, "w2": 50, "w3": 40},
    {"level1": "报考可行性", "level2": "招生路径", "level3": "全日制统考路径明确性", "w1": 30, "w2": 55, "w3": 60},
    {"level1": "报考可行性", "level2": "招生路径", "level3": "名额口径可靠性", "w1": 30, "w2": 55, "w3": 40},
    {"level1": "报考可行性", "level2": "考试风险", "level3": "初复试科目稳定性", "w1": 30, "w2": 45, "w3": 55},
    {"level1": "报考可行性", "level2": "考试风险", "level3": "复试规则可预期性", "w1": 30, "w2": 45, "w3": 45},
    {"level1": "研究培养", "level2": "科研匹配", "level3": "研究方向相关性", "w1": 20, "w2": 55, "w3": 55},
    {"level1": "研究培养", "level2": "科研匹配", "level3": "导师招生路径证据", "w1": 20, "w2": 55, "w3": 45},
    {"level1": "研究培养", "level2": "培养可进入", "level3": "实践机制适用性", "w1": 20, "w2": 45, "w3": 50},
    {"level1": "研究培养", "level2": "培养可进入", "level3": "平台资源可进入性", "w1": 20, "w2": 45, "w3": 50},
    {"level1": "地域生态", "level2": "产业机会", "level3": "城市AI岗位生态", "w1": 10, "w2": 65, "w3": 65},
    {"level1": "地域生态", "level2": "产业机会", "level3": "实习合作半径", "w1": 10, "w2": 65, "w3": 35},
    {"level1": "地域生态", "level2": "生活适配", "level3": "地域与气候", "w1": 10, "w2": 35, "w3": 100},
    {"level1": "成本安排", "level2": "经济成本", "level3": "学费及住宿", "w1": 5, "w2": 55, "w3": 55},
    {"level1": "成本安排", "level2": "经济成本", "level3": "异地与额外支出", "w1": 5, "w2": 55, "w3": 45},
    {"level1": "成本安排", "level2": "培养安排", "level3": "学制与企业实践", "w1": 5, "w2": 45, "w3": 100},
]

# Phase 3.7C 情景倍率（按一级指标名称）
PHASE37C_MULTIPLIERS = {
    "career": {"职业发展": 1.6, "报考可行性": 0.8, "研究培养": 1.0, "地域生态": 1.0, "成本安排": 0.9},
    "admission": {"职业发展": 0.8, "报考可行性": 1.8, "研究培养": 0.9, "地域生态": 0.9, "成本安排": 1.0},
    "research": {"职业发展": 0.9, "报考可行性": 0.8, "研究培养": 1.9, "地域生态": 0.9, "成本安排": 0.9},
}


def make_tree() -> WeightTree:
    return build_tree_from_rows(PHASE37C_ROWS)


def level3_codes(tree: WeightTree) -> list[str]:
    return tree.codes(3)


def code_of(tree: WeightTree, index: int) -> str:
    """按 Excel 行序（0 起）取三级指标代码。"""
    return level3_codes(tree)[index]


def raw_excel_rows() -> list[dict[str, Any]]:
    """直接从原始 Excel 读取指标树（供 golden 校验常量）。"""
    from openpyxl import load_workbook

    wb = load_workbook(RAW_XLSX, data_only=False)
    ws = wb["三级指标树"]
    l1w: dict[str, float] = {}
    l2w: dict[str, float] = {}
    for r in range(2, 7):
        name = ws.cell(r, 12).value
        if name:
            l1w[str(name).strip()] = float(ws.cell(r, 13).value)
    for r in range(9, 19):
        name = ws.cell(r, 12).value
        if name:
            l2w[str(name).strip()] = float(ws.cell(r, 13).value)
    rows: list[dict[str, Any]] = []
    for r in range(2, 20):
        a, b, c = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value
        rows.append(
            {
                "level1": str(a).strip(),
                "level2": str(b).strip(),
                "level3": str(c).strip(),
                "w1": l1w[str(a).strip()],
                "w2": l2w[str(b).strip()],
                "w3": float(ws.cell(r, 6).value),
            }
        )
    return rows
